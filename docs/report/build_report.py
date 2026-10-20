#!/usr/bin/env python3
"""Builds the Final Project Report (.docx) from report.md into the lecturer's template.

Keeps the template cover page, styles and automatic heading numbering; replaces the guidance
body with the report content. Run from the repository root:  python3 docs/report/build_report.py
"""
import copy
import os
import re
import subprocess
import sys
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[2]
TEMPLATE = ROOT / 'docs/report/NPRJ740_Final_Project_Report_Template.docx'
SOURCE = ROOT / 'docs/report/report.md'
OUT = ROOT / 'docs/report/CivicLink_Final_Project_Report.docx'

COVER = {
    5: 'CivicLink: Municipal Incident Reporting and Community Alert Platform',
    7: 'Nthoroane KP',
    8: '202328762',
    17: '30 October 2026',
}


IMG = re.compile(r'!\[(?:Figure:?\s*)?(.*?)\]\((.*?)\)(\{width=([\d.]+)\})?')


def image(path: str) -> Path:
    img = ROOT / path
    if not img.exists():
        sys.exit(f'missing image: {path}')
    return img


def figure_label(caption: str, fig: int) -> tuple[str, int]:
    """Auto-numbers figures; '10a: text' / '10b: text' gives explicit sub-figure labels."""
    m = re.match(r'^(\d+)([a-z]):\s*(.*)', caption)
    if m:
        number = int(m.group(1))
        return f'{number}{m.group(2)}: {m.group(3)}', number
    return f'{fig + 1}: {caption}', fig + 1


def expand_includes(text: str) -> str:
    """{{INCLUDE:path}} inserts a Markdown file; {{TRACE_TABLE}} a condensed traceability table."""
    text = re.sub(r'\{\{INCLUDE:([^}]+)\}\}', lambda m: (ROOT / m.group(1).strip()).read_text(encoding='utf8'), text)
    if '{{TRACE_TABLE}}' in text:
        rows = ['Table: Requirements traceability (condensed from docs/trackers/traceability-matrix.md)', '',
                '| FR | Business rules | Screen | Verified by | Status |', '|---|---|---|---|---|']
        for line in (ROOT / 'docs/trackers/traceability-matrix.md').read_text(encoding='utf8').splitlines():
            c = [x.strip() for x in line.split('|')]
            if len(c) > 9 and re.fullmatch(r'FR-\d\d', c[1]):
                status = c[9].split(' (')[0]
                rows.append(f"| {c[1]} | {c[3]} | {c[5].replace('`', '')} | {c[8]} | {status} |")
        text = text.replace('{{TRACE_TABLE}}', '\n'.join(rows))
    return text


def add_field(doc, anchor, instr: str, entries: list[str], style: str | None = None):
    """Inserts a Word field (TOC/list) before `anchor`, pre-filled with `entries`; Word refreshes it on open."""
    def mk_run(p, kind=None, text=None, instr_text=None):
        r = OxmlElement('w:r')
        if kind:
            fc = OxmlElement('w:fldChar'); fc.set(qn('w:fldCharType'), kind); r.append(fc)
        if instr_text is not None:
            it = OxmlElement('w:instrText'); it.set(qn('xml:space'), 'preserve'); it.text = instr_text; r.append(it)
        if text is not None:
            t = OxmlElement('w:t'); t.set(qn('xml:space'), 'preserve'); t.text = text; r.append(t)
        p._p.append(r)
    entries = entries or ['(Right-click and choose Update Field to build this list.)']
    paras = []
    for e in entries:
        p = doc.add_paragraph(style=style) if style else doc.add_paragraph()
        paras.append(p)
    mk_run(paras[0], 'begin'); mk_run(paras[0], instr_text=f' {instr} '); mk_run(paras[0], 'separate')
    for p, e in zip(paras, entries):
        mk_run(p, text=e)
    mk_run(paras[-1], 'end')
    for p in paras:
        anchor.addprevious(p._p)


def git(*args: str) -> str:
    try:
        return subprocess.check_output(['git', *args], cwd=ROOT, text=True, stderr=subprocess.DEVNULL).strip()
    except subprocess.CalledProcessError:
        return ''


def placeholders() -> dict[str, str]:
    tag = git('describe', '--tags', '--exact-match') or 'v1.0.0'
    commit = git('rev-parse', 'HEAD') or 'unknown'
    remote = git('remote', 'get-url', 'origin')
    m = re.search(r'github\.com[:/](.+?)(\.git)?$', remote)
    repo = f'https://github.com/{m.group(1)}' if m else remote
    app_url = os.environ.get('APP_URL', '(deployed URL pending)')
    return {'{{TAG}}': tag, '{{COMMIT}}': commit, '{{REPO}}': repo, '{{APP_URL}}': app_url}


def num_pr(paragraph):
    ppr = paragraph._p.pPr
    return copy.deepcopy(ppr.numPr) if ppr is not None and ppr.numPr is not None else None


def set_num(paragraph, numpr):
    if numpr is not None:
        paragraph._p.get_or_add_pPr().append(copy.deepcopy(numpr))


def add_runs(paragraph, text: str, size=None):
    """Inline **bold**, `code` and *italic*."""
    for part in re.split(r'(\*\*[^*]+\*\*|`[^`]+`|\*[^*]+\*)', text):
        if not part:
            continue
        if part.startswith('**'):
            run = paragraph.add_run(part[2:-2])
            run.bold = True
        elif part.startswith('`'):
            run = paragraph.add_run(part[1:-1])
            run.font.name = 'Consolas'
            run.font.size = Pt(9)
        elif part.startswith('*') and len(part) > 2:
            run = paragraph.add_run(part[1:-1])
            run.italic = True
        else:
            run = paragraph.add_run(part)
        if size and not part.startswith('`'):
            run.font.size = size


def shade(cell, hex_fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), hex_fill)
    tc_pr.append(shd)


def repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    el = OxmlElement('w:tblHeader')
    el.set(qn('w:val'), 'true')
    tr_pr.append(el)


def main():
    doc = Document(TEMPLATE)
    paras = doc.paragraphs
    h1_num, h2_num, bullet_num = num_pr(paras[21]), num_pr(paras[26]), num_pr(paras[22])
    for name in ('Figure Caption', 'Table Caption'):
        if name not in [st.name for st in doc.styles]:
            st = doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
            st.base_style = doc.styles['Caption']
            st.quick_style = True
    headings: list[str] = []
    figures: list[str] = []
    tables: list[str] = []
    h1 = h2 = 0

    for idx, text in COVER.items():
        p = paras[idx]
        for r in p.runs[1:]:
            r._r.getparent().remove(r._r)
        if p.runs:
            p.runs[0].text = text
        else:
            p.add_run(text)

    # Remove the template guidance body (everything after the cover), keep the section properties.
    body = doc.element.body
    start = paras[19]._p
    remove = False
    for el in list(body):
        if el is start:
            remove = True
        if remove and el.tag != qn('w:sectPr'):
            body.remove(el)

    subs = placeholders()
    text = SOURCE.read_text(encoding='utf8')
    text = re.sub(r'<!--.*?-->', '', text, flags=re.S)
    text = expand_includes(text)
    for k, v in subs.items():
        text = text.replace(k, v)
    lines = text.splitlines()

    fig = 0
    tab = 0
    pending_caption = None
    i = 0
    while i < len(lines):
        line = lines[i].rstrip()
        if not line.strip():
            i += 1
            continue
        if line.startswith('#! '):
            doc.add_paragraph(line[3:].strip(), style='Heading 1')
            headings.append(line[3:].strip())
        elif line.startswith('# '):
            set_num(doc.add_paragraph(line[2:].strip(), style='Heading 1'), h1_num)
            h1 += 1; h2 = 0
            headings.append(f'{h1}  {line[2:].strip()}')
        elif line.startswith('## '):
            set_num(doc.add_paragraph(line[3:].strip(), style='Heading 2'), h2_num)
            h2 += 1
            headings.append(f'    {h1}.{h2}  {line[3:].strip()}')
        elif line.startswith('### '):
            doc.add_paragraph(line[4:].strip(), style='Heading 3')
        elif line.strip() == '\\pagebreak':
            doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        elif line.startswith('> VERIFY:'):
            p = doc.add_paragraph()
            run = p.add_run('TO VERIFY BEFORE SUBMISSION: ' + line[len('> VERIFY:'):].strip())
            run.font.highlight_color = WD_COLOR_INDEX.YELLOW
            run.bold = True
        elif line.startswith('Table: '):
            pending_caption = line[len('Table: '):].strip()
        elif line.startswith('```'):
            code = []
            i += 1
            while i < len(lines) and not lines[i].startswith('```'):
                code.append(lines[i])
                i += 1
            p = doc.add_paragraph()
            run = p.add_run('\n'.join(code))
            run.font.name = 'Consolas'
            run.font.size = Pt(9)
        elif line.startswith('!['):
            group = []
            j = i
            while j < len(lines):
                if not lines[j].strip():
                    j += 1
                    continue
                if not lines[j].startswith('!['):
                    break
                m = IMG.match(lines[j].strip())
                width = float(m.group(4) or 6.5)
                if group and (width > 3.0 or group[0][2] > 3.0):
                    break
                group.append((m.group(1), m.group(2), width))
                j += 1
                if width > 3.0:
                    break
            i = j - 1
            if group[0][2] > 3.0:
                caption, path, width = group[0]
                label, fig = figure_label(caption, fig)
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.keep_with_next = True
                p.add_run().add_picture(str(image(path)), width=Inches(width))
                cap = doc.add_paragraph(f'Figure {label}', style='Figure Caption')
                figures.append(f'Figure {label}')
                cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            else:
                for k in range(0, len(group), 3):
                    chunk = group[k:k + 3]
                    grid = doc.add_table(rows=2, cols=len(chunk))
                    grid.alignment = WD_TABLE_ALIGNMENT.CENTER
                    for c, (caption, path, width) in enumerate(chunk):
                        label, fig = figure_label(caption, fig)
                        pic = grid.cell(0, c).paragraphs[0]
                        pic.alignment = WD_ALIGN_PARAGRAPH.CENTER
                        pic.add_run().add_picture(str(image(path)), width=Inches(min(width, 2.1)))
                        cap = grid.cell(1, c).paragraphs[0]
                        cap.style = doc.styles['Figure Caption']
                        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                        cap.add_run(f'Figure {label}')
                        figures.append(f'Figure {label}')
                    doc.add_paragraph()
        elif line.startswith('|'):
            rows = []
            while i < len(lines) and lines[i].startswith('|'):
                cells = [c.strip() for c in lines[i].strip().strip('|').split('|')]
                if not all(re.fullmatch(r':?-{2,}:?', c) for c in cells):
                    rows.append(cells)
                i += 1
            i -= 1
            if pending_caption:
                tab += 1
                cap = doc.add_paragraph(f'Table {tab}: {pending_caption}', style='Table Caption')
                tables.append(f'Table {tab}: {pending_caption}')
                cap.paragraph_format.keep_with_next = True
                pending_caption = None
            ncols = len(rows[0])
            table = doc.add_table(rows=len(rows), cols=ncols)
            table.style = 'Table Grid'
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            table.autofit = False
            # Each column gets room for its longest word, then the rest is shared by content length.
            char_in = 0.068 if ncols >= 6 else 0.074
            mins, avgs = [], []
            for c in range(ncols):
                col = [re.sub(r'[*`]', '', r[c] if c < len(r) else '') for r in rows]
                longest = max((len(w) for x in col for w in x.split()), default=3)
                mins.append(min(longest * char_in + 0.18, 1.8))
                avgs.append(max(sum(len(x) for x in col) / len(col), 1.0))
            spare = max(6.5 - sum(mins), 0)
            weights = [m + spare * a / sum(avgs) for m, a in zip(mins, avgs)]
            total = sum(weights)
            widths = [Inches(6.5 * w / total) for w in weights]
            for c_idx, col in enumerate(table.columns):
                col.width = widths[c_idx]
            tbl_pr = table._tbl.tblPr
            layout = OxmlElement('w:tblLayout')
            layout.set(qn('w:type'), 'fixed')
            tbl_pr.append(layout)
            size = Pt(8) if ncols >= 6 else Pt(9)
            for r_idx, row in enumerate(rows):
                for c_idx in range(ncols):
                    cell = table.cell(r_idx, c_idx)
                    cell.width = widths[c_idx]
                    cell.paragraphs[0].paragraph_format.space_after = Pt(0)
                    add_runs(cell.paragraphs[0], row[c_idx] if c_idx < len(row) else '', size=size)
                    if r_idx == 0:
                        shade(cell, 'D5EBF4')
                        for run in cell.paragraphs[0].runs:
                            run.bold = True
            repeat_header(table.rows[0])
            doc.add_paragraph()
        elif line.startswith('- '):
            p = doc.add_paragraph(style='Bullet Compact')
            set_num(p, bullet_num)
            add_runs(p, line[2:].strip())
        elif re.match(r'^\d+\. ', line):
            p = doc.add_paragraph(style='List Number')
            add_runs(p, re.sub(r'^\d+\. ', '', line).strip())
        else:
            paragraph = [line.strip()]
            while i + 1 < len(lines) and lines[i + 1].strip() and not re.match(r'^(#|- |\||!\[|Table: |> |```|\\pagebreak|\d+\. )', lines[i + 1]):
                i += 1
                paragraph.append(lines[i].strip())
            p = doc.add_paragraph()
            add_runs(p, ' '.join(paragraph))
        i += 1

    # Front matter: contents, figures and tables, inserted before the Executive Summary.
    first = next(p for p in doc.paragraphs if p.style.name == 'Heading 1')._p
    for title, instr, entries in (
        ('Table of Contents', 'TOC \\o "1-2" \\h \\z \\u', headings),
        ('List of Figures', 'TOC \\h \\z \\t "Figure Caption,1"', figures),
        ('List of Tables', 'TOC \\h \\z \\t "Table Caption,1"', tables),
    ):
        h = doc.add_paragraph(title, style='TOC Heading')
        first.addprevious(h._p)
        add_field(doc, first, instr, entries)
        br = doc.add_paragraph()
        br.add_run().add_break(WD_BREAK.PAGE)
        first.addprevious(br._p)
    # Ask Word to refresh the fields (adds page numbers) when the document is opened.
    upd = OxmlElement('w:updateFields'); upd.set(qn('w:val'), 'true')
    doc.settings.element.append(upd)
    # Visible heading colour consistency is inherited from the template styles.
    _ = RGBColor
    doc.save(OUT)
    print(f'Wrote {OUT.relative_to(ROOT)}: {fig} figures, {tab} tables; tag {subs["{{TAG}}"]}, commit {subs["{{COMMIT}}"][:7]}')


if __name__ == '__main__':
    main()
